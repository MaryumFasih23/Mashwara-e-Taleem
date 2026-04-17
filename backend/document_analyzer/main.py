import os, re, json, pickle, tempfile, shutil, asyncio
import numpy as np
import pandas as pd
from pypdf import PdfReader
import docx
import language_tool_python
from sentence_transformers import SentenceTransformer
import faiss
from groq import Groq
from fastapi import FastAPI, UploadFile, File, Form
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse
from contextlib import asynccontextmanager
from dotenv import load_dotenv

load_dotenv()

# ── CONFIG ──────────────────────────────────────────────────────
GROQ_API_KEY = os.environ.get("GROQ_API_KEY")
LLM_MODEL    = "llama-3.3-70b-versatile"
EMBED_MODEL  = "sentence-transformers/all-MiniLM-L6-v2"

DATA_DIR = os.path.join(os.path.dirname(__file__), "data")

US_UNI_CSV      = os.path.join(DATA_DIR, "us_universities.csv")
US_PROG_CSV     = os.path.join(DATA_DIR, "us_programs.csv")
NON_US_UNI_CSV  = os.path.join(DATA_DIR, "non_us_universities.csv")
NON_US_PROG_CSV = os.path.join(DATA_DIR, "non_us_programs.csv")

UNI_EMB_PKL  = os.path.join(DATA_DIR, "uni_embeddings.pkl")
PROG_EMB_PKL = os.path.join(DATA_DIR, "prog_embeddings.pkl")

KEYWORD_STOPLIST = {
    "the","and","for","with","this","that","from","your","have","will","you","our","are","can","also",
    "program","university","department","students","course","courses","study","research","faculty",
    "requirements","requirement","application","admission","deadline","tuition","fees","location",
    "more","information","please","visit","contact","email","phone","address","apply","online",
    "full","time","part","per","year","credit","hours","minimum","maximum","based","must","meet",
    "all","any","one","two","three","four","five","six","seven","eight","nine","ten",
    "not","but","its","has","was","been","they","their","them","who","what","when","where","how",
    "each","which","about","after","before","during","other","into","than","then","some","only",
    "may","use","used","new","also","both","well","even","such","same","most","many","best",
    "http","https","www","html","php","asp","aspx","edu","com","org","net","gov","htm",
    "nan","none","null","true","false","yes","no",
    "rolling","adance","january","advancedcomputinginstitute","institute","private","college",
    "academics","programs","computing","public",
}

ACTION_VERBS = [
    "led","developed","implemented","designed","built","optimized","created",
    "managed","analyzed","improved","increased","reduced","delivered","automated",
    "coordinated","executed","launched","engineered","refactored","deployed","integrated"
]

TECHNICAL_TERMS = {
    "FastAPI", "WhisperX", "LangChain", "Prisma", "FAISS", "NaSCon", "ICAP", "IBA",
    "WordPress", "PostgreSQL", "PyTorch", "CLIP", "WebRTC", "Firebase", "Puppeteer",
    "Tailwind", "TypeScript", "JavaScript", "Django", "React", "Node.js", "MySQL",
    "LLMs", "Tetris", "tetrominoes", "KAIRO", "Kairo", "TeachTrack", "Bootcamp",
}
TECHNICAL_TERMS_LOWER = {t.lower() for t in TECHNICAL_TERMS}

COMMON_CAPITALIZED_WORDS = {
    "a", "an", "and", "as", "at", "but", "by", "for", "from", "if", "in", "into",
    "is", "it", "of", "on", "or", "the", "to", "with", "i",
    "january", "february", "march", "april", "may", "june", "july", "august",
    "september", "october", "november", "december",
}

INFORMAL_WORD_REPLACEMENTS = {
    "stuff": "work",
    "things": "details",
    "a lot of": "many",
    "lots of": "many",
    "really": "",
    "very": "",
    "pretty": "",
    "got": "received",
    "gonna": "going to",
    "wanna": "want to",
    "kinda": "somewhat",
    "sort of": "somewhat",
}

ARTICLE_NOUNS = (
    "student", "developer", "engineer", "intern", "assistant", "researcher",
    "candidate", "graduate", "undergraduate", "officer", "member", "leader",
    "application", "app", "system", "platform", "tool", "model", "database",
    "dashboard", "website", "api", "project", "course", "program",
)

DOMAIN_KEYWORDS = {
    "computer_science": {
        "computer", "science", "software", "programming", "algorithm", "algorithms",
        "data", "database", "ai", "artificial", "intelligence", "machine", "learning",
        "python", "java", "javascript", "typescript", "react", "node", "fastapi",
        "api", "backend", "frontend", "cloud", "security", "network", "systems",
        "engineering", "developer", "web", "application", "model", "models",
        "faiss", "langchain", "prisma", "postgresql", "mysql", "firebase",
    },
    "law": {
        "law", "legal", "court", "courts", "case", "cases", "contract", "contracts",
        "constitution", "constitutional", "policy", "rights", "justice", "litigation",
        "advocacy", "regulation", "regulatory", "compliance", "criminal", "civil",
        "corporate", "jurisprudence", "legislation", "statute", "statutes", "trial",
        "moot", "clerkship", "attorney", "lawyer", "evidence", "ethics",
    },
    "business": {
        "business", "management", "finance", "financial", "marketing", "market",
        "markets", "strategy", "strategic", "operations", "entrepreneurship",
        "startup", "revenue", "sales", "accounting", "investment", "investments",
        "consulting", "analytics", "economics", "leadership", "supply", "chain",
        "product", "customer", "customers", "brand", "profit", "budget", "growth",
        "mba", "commerce",
    },
    "arts": {
        "arts", "art", "design", "visual", "creative", "creativity", "studio",
        "painting", "drawing", "sculpture", "illustration", "photography",
        "film", "media", "animation", "portfolio", "gallery", "museum",
        "history", "culture", "cultural", "humanities", "literature", "music",
        "theatre", "theater", "performance", "aesthetic", "composition",
    },
}

DOMAIN_LABELS = {
    "computer_science": "Computer Science",
    "law": "Law",
    "business": "Business",
    "arts": "Arts",
}

DOMAIN_KEYWORD_MAPPING = {
    "ai": {"ai", "artificial intelligence", "machine learning", "deep learning", "nlp", "natural language processing"},
    "backend": {"backend", "fastapi", "node", "node.js", "api", "apis", "server", "server-side"},
    "frontend": {"frontend", "front-end", "react", "next.js", "nextjs", "ui", "user interface"},
    "data": {"data", "sql", "database", "databases", "mongodb", "mongo", "postgresql", "mysql"},
    "cloud": {"cloud", "firebase", "supabase", "vercel", "deployment", "hosting"},
}

# ── GLOBAL STATE (loaded once at startup) ───────────────────────
_state = {
    "client":     None,
    "embedder":   None,
    "lt_tool":    None,
    "unis":       None,
    "progs":      None,
    "uni_index":  None,
    "prog_index": None,
}


# ── DATASET NORMALIZATION ────────────────────────────────────────

def normalize_us_universities(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    out["university_name"] = out.get("University", pd.Series(dtype=str))
    out["program_name"]    = ""
    return out

def normalize_non_us_universities(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    out["program_name"] = ""
    return out

def normalize_us_programs(df: pd.DataFrame) -> pd.DataFrame:
    return df.copy()

def normalize_non_us_programs(df: pd.DataFrame) -> pd.DataFrame:
    return df.copy()


def load_data():
    frames_uni  = []
    frames_prog = []

    if os.path.exists(US_UNI_CSV):
        df = pd.read_csv(US_UNI_CSV)
        df = normalize_us_universities(df)
        frames_uni.append(df)

    if os.path.exists(NON_US_UNI_CSV):
        df = pd.read_csv(NON_US_UNI_CSV)
        df = normalize_non_us_universities(df)
        frames_uni.append(df)

    if os.path.exists(US_PROG_CSV):
        df = pd.read_csv(US_PROG_CSV)
        df = normalize_us_programs(df)
        frames_prog.append(df)

    if os.path.exists(NON_US_PROG_CSV):
        df = pd.read_csv(NON_US_PROG_CSV)
        df = normalize_non_us_programs(df)
        frames_prog.append(df)

    unis  = pd.concat(frames_uni,  ignore_index=True) if frames_uni  else pd.DataFrame()
    progs = pd.concat(frames_prog, ignore_index=True) if frames_prog else pd.DataFrame()

    unis["combined"]  = unis.fillna("").astype(str).agg(" | ".join, axis=1)
    progs["combined"] = progs.fillna("").astype(str).agg(" | ".join, axis=1)

    return unis, progs


# ── EMBEDDING + FAISS ────────────────────────────────────────────

def build_index(df: pd.DataFrame, text_col: str, pkl_path: str, embedder):
    texts = df[text_col].fillna("").astype(str).tolist()
    if os.path.exists(pkl_path):
        with open(pkl_path, "rb") as f:
            emb = pickle.load(f)
        emb = np.array(emb, dtype="float32")
    else:
        emb = embedder.encode(texts, normalize_embeddings=True, show_progress_bar=True)
        emb = np.array(emb, dtype="float32")
        with open(pkl_path, "wb") as f:
            pickle.dump(emb, f)
    index = faiss.IndexFlatIP(emb.shape[1])
    index.add(emb)
    return index

def retrieve(df: pd.DataFrame, index, query: str, embedder, k: int = 3) -> pd.DataFrame:
    q = embedder.encode([query], normalize_embeddings=True).astype("float32")
    D, I = index.search(q, k)
    out = df.iloc[I[0]].copy()
    out["score"] = D[0]
    return out


def retrieve_from_embedding(df: pd.DataFrame, index, query_embedding, k: int = 3) -> pd.DataFrame:
    D, I = index.search(query_embedding, k)
    out = df.iloc[I[0]].copy()
    out["score"] = D[0]
    return out


def retrieve_context(unis, progs, uni_index, prog_index, query: str, embedder, k: int = 3) -> str:
    query_embedding = embedder.encode([query], normalize_embeddings=True).astype("float32")
    u_rows = retrieve_from_embedding(unis, uni_index, query_embedding, k=k)
    p_rows = retrieve_from_embedding(progs, prog_index, query_embedding, k=k)
    return "\n".join(list(u_rows["combined"]) + list(p_rows["combined"]))


# ── LIFESPAN (replaces deprecated @app.on_event) ────────────────

@asynccontextmanager
async def lifespan(app: FastAPI):
    """Load all heavy resources once at startup."""
    print("[startup] Loading Groq client...")
    _state["client"] = Groq(api_key=GROQ_API_KEY)

    print("[startup] Loading SentenceTransformer...")
    _state["embedder"] = SentenceTransformer(EMBED_MODEL)

    print("[startup] Loading LanguageTool (requires Java)...")
    _state["lt_tool"] = language_tool_python.LanguageTool("en-US")

    print("[startup] Loading CSV datasets...")
    unis, progs = load_data()
    _state["unis"]  = unis
    _state["progs"] = progs

    print("[startup] Building FAISS indexes...")
    _state["uni_index"]  = build_index(unis,  "combined", UNI_EMB_PKL,  _state["embedder"])
    _state["prog_index"] = build_index(progs, "combined", PROG_EMB_PKL, _state["embedder"])

    print("[startup] All resources loaded. Server ready.")
    yield

    # Cleanup on shutdown
    if _state["lt_tool"]:
        _state["lt_tool"].close()
    print("[shutdown] Resources released.")


app = FastAPI(lifespan=lifespan)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)


# ── HELPERS ─────────────────────────────────────────────────────

def safe_json_parse(text):
    if text is None:
        return {"error": "empty_response"}
    t = text.strip()
    if t.startswith("```"):
        t = re.sub(r"^```[a-zA-Z]*\s*", "", t)
        t = re.sub(r"\s*```$", "", t).strip()
    try:
        return json.loads(t)
    except Exception:
        m = re.search(r"\{.*\}", t, re.DOTALL)
        if m:
            try:
                return json.loads(m.group(0))
            except Exception:
                pass
    return {"error": "invalid_json", "raw": text[:4000]}

def ask_llm_json(prompt, max_tokens=2500, temperature=0.0):
    client = _state["client"]
    r = client.chat.completions.create(
        model=LLM_MODEL,
        messages=[{"role": "user", "content": prompt}],
        temperature=temperature,
        max_tokens=max_tokens
    )
    return safe_json_parse(r.choices[0].message.content)

def read_pdf(path):
    reader = PdfReader(path)
    return "\n".join([(p.extract_text() or "") for p in reader.pages]).strip()

def read_docx(path):
    d = docx.Document(path)
    return "\n".join([p.text for p in d.paragraphs]).strip()

def read_txt(path):
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        return f.read().strip()

def read_document(path):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":  return read_pdf(path)
    if ext == ".docx": return read_docx(path)
    if ext == ".txt":  return read_txt(path)
    raise ValueError("Unsupported file type. Please upload .pdf, .docx, or .txt")


def _looks_like_address(line: str) -> bool:
    return bool(re.search(r"\b(street|st\.|road|rd\.|avenue|ave\.|lane|ln\.|city|zip|postal|block|house|apartment|apt)\b", line, re.I))


def _looks_like_name_only_line(stripped: str) -> bool:
    """Whole line is likely a person's name (2–5 title-case tokens, no digits)."""
    if not stripped or len(stripped) > 120:
        return False
    if re.search(r"\d|@|\+|/", stripped):
        return False
    return bool(re.fullmatch(r"[A-Z][a-z]+(?:\s+[A-Z][a-z]+){1,4}\s*", stripped))


def _looks_like_contact_header_line(stripped: str) -> bool:
    """Resume header lines: roll no, phone, email, LinkedIn, etc."""
    low = stripped.lower()
    if "@" in stripped:
        return True
    if re.search(r"\+\d{1,3}[\s\-]?\d{3,}", stripped):
        return True
    if re.search(r"\b(roll\s*no|linkedin|github|phone|mobile|email)\b", low):
        return True
    # Student / roll IDs like 22i-1244, 22I-1244
    if re.search(r"\b\d{1,3}[a-zA-Z][-]?\d{3,}\b", stripped):
        return True
    if re.search(r"\b\d{2}[a-z]\s*[-–]?\s*\d{4,}\b", low):
        return True
    return False


def _lt_rule_id(m) -> str:
    return str(getattr(m, "rule_id", None) or getattr(m, "ruleId", "") or "")


def _lt_category(m) -> str:
    cat = getattr(m, "category", "") or ""
    return str(getattr(cat, "name", cat) or "")


def _clean_word_token(token: str) -> str:
    return re.sub(r"^[^A-Za-z0-9+#.]+|[^A-Za-z0-9+#.]+$", "", token or "")


def _is_protected_term_token(token: str) -> bool:
    clean = _clean_word_token(token)
    if not clean:
        return False
    if clean.lower() in TECHNICAL_TERMS_LOWER:
        return True
    if any(t.lower() == clean.lower() for t in TECHNICAL_TERMS):
        return True
    if re.search(r"[A-Za-z]\d|\d[A-Za-z]", clean):
        return True
    if clean.isupper() and len(clean) >= 2:
        return True
    if re.search(r"[a-z][A-Z]|[A-Z][a-z]+[A-Z]", clean):
        return True
    if clean[0].isupper() and clean.lower() not in COMMON_CAPITALIZED_WORDS:
        return True
    return False


def _is_protected_term_text(text: str) -> bool:
    clean = _clean_word_token(text)
    return bool(clean and _is_protected_term_token(clean))


def _line_has_protected_term_changed(original: str, improved: str) -> bool:
    if not original or not improved:
        return False
    protected = [
        t for t in re.findall(r"\b[A-Za-z][A-Za-z0-9+#.]*\b", original)
        if _is_protected_term_token(t)
    ]
    for token in protected:
        if token not in improved:
            return True
    return False


def _is_spelling_match(m) -> bool:
    rid = _lt_rule_id(m).upper()
    category = _lt_category(m).upper()
    msg = (getattr(m, "message", "") or "").lower()
    return (
        "SPELL" in rid
        or "MORFOLOGIK" in rid
        or "HUNSPELL" in rid
        or "TYPOS" in category
        or "spelling" in msg
    )


def _edit_distance(a: str, b: str) -> int:
    if a == b:
        return 0
    if not a:
        return len(b)
    if not b:
        return len(a)
    prev = list(range(len(b) + 1))
    for i, ca in enumerate(a, start=1):
        cur = [i]
        for j, cb in enumerate(b, start=1):
            cur.append(min(
                prev[j] + 1,
                cur[j - 1] + 1,
                prev[j - 1] + (0 if ca == cb else 1),
            ))
        prev = cur
    return prev[-1]


def _spelling_suggestion_high_confidence(original: str, replacement: str) -> bool:
    src = _clean_word_token(original)
    rep = _clean_word_token(replacement)
    if not src or not rep:
        return False
    if _is_protected_term_token(src) or _is_protected_term_token(rep):
        return False
    if not re.fullmatch(r"[a-z]+", src) or not re.fullmatch(r"[a-z]+", rep):
        return False
    if src[0] != rep[0]:
        return False
    dist = _edit_distance(src, rep)
    limit = 1 if len(src) <= 7 else 2
    return dist <= limit


def _should_keep_spelling_match(m, matched_text: str) -> bool:
    if _is_protected_term_text(matched_text):
        return False
    reps = getattr(m, "replacements", []) or []
    if len(reps) != 1:
        return False
    rep = _first_replacement_str(reps)
    return _spelling_suggestion_high_confidence(matched_text, rep)


def _classify_lt_issue(m) -> str:
    rid = _lt_rule_id(m).upper()
    category = _lt_category(m).upper()
    msg = (getattr(m, "message", "") or "").lower()
    if "PUNCT" in rid or "TYPOGRAPHY" in category or "dash" in msg or "hyphen" in msg:
        return "formatting"
    if "CASING" in rid or "CASE" in rid or "uppercase" in msg or "lowercase" in msg:
        return "formatting"
    if "STYLE" in category or "READABILITY" in category or "wordy" in msg or "simpl" in msg:
        return "clarity"
    return "grammar"


def _should_skip_lt_match(m, text: str) -> bool:
    """
    Drop noisy LanguageTool findings that are not useful for admissions documents:
    double spaces, trivial whitespace, etc.
    """
    rid = _lt_rule_id(m).upper()
    msg = (getattr(m, "message", "") or "").lower() if m else ""
    # Whitespace-only / duplicate-space rules
    if "WHITESPACE" in rid or ("DUPLICATE" in rid and "SPACE" in rid):
        return True
    if "whitespace" in msg and ("repeat" in msg or "duplicate" in msg or "twice" in msg):
        return True
    if "repeated" in msg and "whitespace" in msg:
        return True
    # Often flags "  " between label and value — low value for reviewers
    if "two consecutive" in msg and "space" in msg:
        return True
    return False


def _find_protected_spans(text: str):
    """
    Protect personal identifiers from auto-correction to avoid damaging
    user-provided facts like names, contact information, tools, acronyms,
    and proper nouns.
    """
    spans = []

    # Emails
    for m in re.finditer(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}", text):
        spans.append((m.start(), m.end()))

    # Phone numbers (simple international/local patterns)
    for m in re.finditer(r"(?:(?:\+?\d{1,3}[\s\-]?)?(?:\(?\d{2,4}\)?[\s\-]?)?\d{3,4}[\s\-]?\d{3,4})", text):
        if len(re.sub(r"\D", "", m.group(0))) >= 7:
            spans.append((m.start(), m.end()))

    # Whole-line protection: name-only lines, contact headers, addresses (exact offsets)
    pos = 0
    for raw in text.splitlines(keepends=True):
        line = raw.rstrip("\r\n")
        stripped = line.strip()
        if stripped:
            line_start = pos
            end = pos + len(line)
            if (
                _looks_like_name_only_line(stripped)
                or _looks_like_contact_header_line(stripped)
                or _looks_like_address(line)
            ):
                spans.append((line_start, end))
        pos += len(raw)

    # Capitalized multi-token name chunks anywhere in text
    for m in re.finditer(r"\b[A-Z][a-z]{1,20}(?:\s+[A-Z][a-z]{1,20}){1,3}\b", text):
        phrase = m.group(0)
        if not _looks_like_address(phrase) and len(phrase.split()) <= 4:
            spans.append((m.start(), m.end()))

    # Exact technical/product terms and acronym/camel-case/proper-noun tokens.
    for term in sorted(TECHNICAL_TERMS, key=len, reverse=True):
        pattern = r"(?<![A-Za-z0-9])" + re.escape(term) + r"(?![A-Za-z0-9])"
        for m in re.finditer(pattern, text):
            spans.append((m.start(), m.end()))
    for m in re.finditer(r"\b[A-Za-z][A-Za-z0-9+#.]*\b", text):
        if _is_protected_term_token(m.group(0)):
            spans.append((m.start(), m.end()))

    return spans


def _overlaps_protected(start: int, end: int, spans):
    for s, e in spans:
        if start < e and end > s:
            return True
    return False


def _extract_protected_values(text: str):
    values = set()
    for term in TECHNICAL_TERMS:
        if re.search(r"(?<![A-Za-z0-9])" + re.escape(term) + r"(?![A-Za-z0-9])", text):
            values.add(term)
    for m in re.finditer(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}", text):
        values.add(m.group(0))
    for m in re.finditer(r"(?:(?:\+?\d{1,3}[\s\-]?)?(?:\(?\d{2,4}\)?[\s\-]?)?\d{3,4}[\s\-]?\d{3,4})", text):
        if len(re.sub(r"\D", "", m.group(0))) >= 7:
            values.add(m.group(0).strip())
    lines = text.splitlines()
    if lines:
        first = lines[0].strip()
        if re.fullmatch(r"[A-Z][a-z]+(?:\s+[A-Z][a-z]+){1,3}", first):
            values.add(first)
    for m in re.finditer(r"\b[A-Z][a-z]{1,20}(?:\s+[A-Z][a-z]{1,20}){1,3}\b", text):
        phrase = m.group(0)
        if not _looks_like_address(phrase):
            values.add(phrase)
    for m in re.finditer(r"\b[A-Za-z][A-Za-z0-9+#.]*\b", text):
        token = m.group(0)
        if _is_protected_term_token(token):
            values.add(token)
    return sorted(values, key=len, reverse=True)


def _line_no_from_offset(text: str, offset: int):
    return text[:max(0, offset)].count("\n") + 1


def _normalize_ws(s: str) -> str:
    return re.sub(r"\s+", " ", (s or "").strip())


def _get_line_at_number(text: str, line_no: int) -> str:
    lines = text.splitlines()
    if 1 <= line_no <= len(lines):
        return lines[line_no - 1]
    return ""


def _error_len(m) -> int:
    return int(getattr(m, "errorLength", None) or getattr(m, "error_length", 0) or 0)


def _looks_like_section_heading(stripped: str) -> bool:
    """Short resume section titles (Education, PROJECTS, etc.) — not grammar errors."""
    if not stripped or len(stripped) > 72:
        return False
    low = stripped.lower().strip(":：")
    common = (
        "education", "experience", "projects", "skills", "summary", "objective",
        "contact", "work", "employment", "certifications", "awards", "additional",
        "references", "publications", "leadership", "activities", "coursework",
    )
    if low in common or low.rstrip("s") in common:
        return True
    if re.fullmatch(r"[A-Za-z][A-Za-z\s/&,.-]{0,50}", stripped) and len(stripped.split()) <= 5:
        if stripped.isupper() or stripped.istitle():
            return True
    return False


def _is_case_or_punct_only_change(o: str, i: str) -> bool:
    return _normalize_ws(o).lower() == _normalize_ws(i).lower()


def _sanitize_placeholder_text(s: str) -> str:
    """Replace vague [ADD DETAIL] style placeholders with actionable hints."""
    if not s:
        return s
    s = re.sub(r"\[ADD DETAIL\]", "a specific metric or outcome (e.g., latency reduced by 20%, 500+ users)", s, flags=re.I)
    s = re.sub(r"\[ADD\s+METRIC\]", "a measurable result (%, time saved, scale)", s, flags=re.I)
    return s


def _grammar_score_from_count(issue_count: int) -> int:
    """
    Score only true grammar errors. A clean or nearly clean SOP should stay high,
    while drafts with many grammar errors should drop quickly.
    """
    n = max(0, int(issue_count or 0))
    if n == 0:
        return 100
    if n <= 3:
        return max(0, 100 - n * 3)
    if n <= 8:
        return max(0, 91 - (n - 3) * 4)
    return max(0, 71 - (n - 8) * 5)


def _is_real_grammar_issue(issue: dict) -> bool:
    return isinstance(issue, dict) and str(issue.get("issue_type", "")).lower() == "grammar"


def _grammar_error_count(issues: list) -> int:
    return sum(1 for issue in issues if _is_real_grammar_issue(issue))


def _article_for_word(word: str) -> str:
    return "an" if re.match(r"(?i)[aeiou]", word or "") else "a"


def _is_real_sentence_line(stripped: str) -> bool:
    if not stripped or _looks_like_section_heading(stripped):
        return False
    if _looks_like_name_only_line(stripped) or _looks_like_contact_header_line(stripped):
        return False
    if re.match(r"^(https?://|www\.|[\w.%+-]+@)", stripped, re.I):
        return False
    if len(stripped.split()) < 4:
        return False
    return bool(re.search(r"[.!?]$", stripped) or re.search(r"\b(is|are|am|was|were|have|has|had|built|developed|worked|created|designed|implemented|led|managed|did|made)\b", stripped, re.I))


def _make_heuristic_issue(line_no: int, issue_type: str, severity: str, original: str, improved: str, explanation: str, rule_id: str) -> dict:
    return {
        "message": explanation,
        "context": original,
        "replacements": [improved],
        "offset": 0,
        "error_length": len(original),
        "line_number": line_no,
        "original_line": original,
        "issue_type": issue_type,
        "severity": severity,
        "snippet": original,
        "source": "heuristic",
        "rule_id": rule_id,
    }


def _detect_lowercase_sentence_starts(line_no: int, line: str) -> list:
    stripped = line.strip()
    if not _is_real_sentence_line(stripped):
        return []
    issues = []
    first_alpha = re.search(r"[A-Za-z]", line)
    if first_alpha and line[first_alpha.start()].islower():
        improved = line[:first_alpha.start()] + line[first_alpha.start()].upper() + line[first_alpha.start() + 1:]
        issues.append(_make_heuristic_issue(
            line_no, "grammar", "minor", line, improved,
            "Sentence starts should be capitalized. Section headings are ignored for this check.",
            "HEURISTIC_LOWERCASE_SENTENCE_START",
        ))
    for m in re.finditer(r"([.!?]\s+)([a-z])", line):
        improved = line[:m.start(2)] + line[m.start(2)].upper() + line[m.start(2) + 1:]
        issues.append(_make_heuristic_issue(
            line_no, "grammar", "minor", line, improved,
            "A new sentence appears to start with a lowercase letter.",
            "HEURISTIC_LOWERCASE_SENTENCE_START",
        ))
    return issues


def _detect_informal_words(line_no: int, line: str) -> list:
    if _looks_like_section_heading(line.strip()) or _looks_like_contact_header_line(line.strip()):
        return []
    issues = []
    for informal, replacement in INFORMAL_WORD_REPLACEMENTS.items():
        pattern = r"\b" + re.escape(informal) + r"\b"
        m = re.search(pattern, line, re.I)
        if not m:
            continue
        if replacement:
            improved = line[:m.start()] + replacement + line[m.end():]
            improved = re.sub(r"\s{2,}", " ", improved).strip()
        else:
            improved = (line[:m.start()] + line[m.end():]).replace("  ", " ").strip()
        issues.append(_make_heuristic_issue(
            line_no, "clarity", "minor", line, improved,
            f"'{m.group(0)}' is informal; use more precise academic or resume wording.",
            "HEURISTIC_INFORMAL_WORD",
        ))
    return issues


def _detect_missing_articles(line_no: int, line: str) -> list:
    if _looks_like_section_heading(line.strip()) or _looks_like_contact_header_line(line.strip()):
        return []
    issues = []
    noun_pattern = "|".join(re.escape(n) for n in ARTICLE_NOUNS)
    be_pattern = re.compile(
        rf"\b(am|is|are|was|were|be|been|being|became|become)\s+((?!(?:a|an|the|my|our|your|their|this|that)\b)(?:[a-z][a-z+-]*\s+){{0,4}})({noun_pattern})\b",
        re.I,
    )
    action_pattern = re.compile(
        rf"\b(built|developed|created|designed|implemented|made|launched|trained|deployed)\s+((?!(?:a|an|the|my|our|your|their|this|that)\b)(?:[a-z][a-z+-]*\s+){{0,3}})({noun_pattern})\b",
        re.I,
    )
    for pattern in (be_pattern, action_pattern):
        m = pattern.search(line)
        if not m:
            continue
        phrase_start = m.start(2)
        phrase = (m.group(2) or "") + (m.group(3) or "")
        first_word = (phrase.split() or [m.group(3)])[0]
        article = _article_for_word(first_word)
        improved = line[:phrase_start] + article + " " + line[phrase_start:]
        if _normalize_ws(improved).lower() == _normalize_ws(line).lower():
            continue
        issues.append(_make_heuristic_issue(
            line_no, "grammar", "important", line, improved,
            f"Possible missing article before '{phrase.strip()}'.",
            "HEURISTIC_MISSING_ARTICLE",
        ))
    return issues


def _detect_weak_sentence_structure(line_no: int, line: str) -> list:
    stripped = line.strip()
    if not stripped or _looks_like_section_heading(stripped) or _looks_like_contact_header_line(stripped):
        return []
    issues = []
    words = re.findall(r"[A-Za-z0-9+#.]+", stripped)
    and_count = len(re.findall(r"\band\b", stripped, re.I))
    if len(words) >= 42 or and_count >= 4:
        improved = "Split into two shorter sentences: " + stripped
        issues.append(_make_heuristic_issue(
            line_no, "clarity", "important", line, improved,
            "This sentence is overloaded; split it or make the structure more direct.",
            "HEURISTIC_WEAK_STRUCTURE_LONG",
        ))
    weak_openers = (
        (r"^\s*worked on\b", "Developed or contributed to"),
        (r"^\s*responsible for\b", "Led"),
        (r"^\s*did\b", "Completed"),
        (r"^\s*made\b", "Built"),
    )
    for pattern, replacement in weak_openers:
        if re.search(pattern, stripped, re.I):
            improved = re.sub(pattern, replacement, line, count=1, flags=re.I)
            issues.append(_make_heuristic_issue(
                line_no, "clarity", "minor", line, improved,
                "Weak sentence structure; start with a stronger, more specific action.",
                "HEURISTIC_WEAK_STRUCTURE_OPENING",
            ))
            break
    return issues


def _heuristic_grammar_issues(text: str) -> list:
    issues = []
    for line_no, _, line in _line_ranges(text):
        issues.extend(_detect_lowercase_sentence_starts(line_no, line))
        issues.extend(_detect_missing_articles(line_no, line))
        issues.extend(_detect_informal_words(line_no, line))
        issues.extend(_detect_weak_sentence_structure(line_no, line))
    return _dedupe_line_issues([i for i in issues if not _is_noise_line_issue(i)])


def _is_llm_line_issue_junk(li: dict, text: str) -> bool:
    """Drop LLM rows that hallucinate names or boilerplate header rewrites."""
    o = (li.get("original_line") or "").strip()
    imp = (li.get("improved_line") or "").strip()
    expl = (li.get("explanation") or "").lower()
    low_imp = imp.lower()
    if "name + contact" in low_imp or "name+contact" in low_imp.replace(" ", ""):
        return True
    if "possible spelling" in expl:
        return True
    if _line_has_protected_term_changed(o, imp):
        return True
    if _looks_like_name_only_line(o) or _looks_like_contact_header_line(o):
        return True
    if _looks_like_section_heading(o) and _is_case_or_punct_only_change(o, imp):
        return True
    # Corrupted name patterns (Area vs Areeba) when line is name-like
    if _looks_like_name_only_line(o) and len(o.split()) <= 4:
        if "area " in low_imp and "areeba" not in low_imp:
            return True
    return False


def _is_noise_line_issue(li: dict) -> bool:
    """Drop no-op or low-value line issues (duplicate whitespace, name false positives, etc.)."""
    o = (li.get("original_line") or "").strip()
    i = (li.get("improved_line") or _first_replacement_str(li.get("replacements") or []) or "").strip()
    expl = (li.get("explanation") or "").lower()
    if not o:
        return True
    if not i or o == i:
        return True
    if _normalize_ws(o).lower() == _normalize_ws(i).lower() and li.get("rule_id") != "HEURISTIC_LOWERCASE_SENTENCE_START":
        return True
    if _normalize_ws(o) == _normalize_ws(i):
        return True
    if "whitespace" in expl or ("repeated" in expl and "space" in expl):
        return True
    if "two consecutive" in expl and "space" in expl:
        return True
    if "possible spelling" in expl:
        return True
    if _line_has_protected_term_changed(o, i):
        return True
    if "spelling" in expl and _looks_like_name_only_line(o):
        return True
    if "spelling" in expl and o.replace(" ", "").isalpha() and _looks_like_name_only_line(o):
        return True
    # Bad suggestions like "F.Sc ... -> Sc"
    if "->" in i and i.split("->")[-1].strip().lower() in ("sc", "f.sc", "f"):
        return True
    return False


def _line_issue_compare_text(value: str) -> str:
    text = _sanitize_placeholder_text(value or "")
    text = re.sub(r"^\s*(original|suggested|improved)\s*:\s*", "", text, flags=re.I)
    text = text.strip().strip("\"'`")
    text = _normalize_ws(text).lower()
    text = re.sub(r"\s+([,.;:!?])", r"\1", text)
    return text


def _is_noop_line_suggestion(original: str, improved: str, rule_id: str = "") -> bool:
    o = _line_issue_compare_text(original)
    i = _line_issue_compare_text(improved)
    if not o or not i:
        return True
    if rule_id == "HEURISTIC_LOWERCASE_SENTENCE_START":
        return False
    if o == i:
        return True
    if re.sub(r"[^a-z0-9]+", "", o) == re.sub(r"[^a-z0-9]+", "", i):
        return True
    if i.startswith(o) and len(i) <= len(o) + 3:
        return True
    return False


def _line_ranges(text: str):
    out = []
    pos = 0
    for line_no, raw in enumerate(text.splitlines(keepends=True), start=1):
        line = raw.rstrip("\r\n")
        out.append((line_no, pos, line))
        pos += len(raw)
    return out


def _first_replacement_str(reps) -> str:
    if not reps:
        return ""
    r = reps[0]
    if isinstance(r, str):
        return r
    return str(getattr(r, "value", None) or r)


def _build_grammar_line_issue(text: str, gi: dict) -> dict:
    """One row from LanguageTool issue with a concrete line-level suggestion."""
    offset = int(gi.get("offset", 0))
    elen = int(gi.get("error_length", 0) or 0)
    if elen <= 0:
        elen = 1
    reps = gi.get("replacements") or []
    rep = _first_replacement_str(reps)
    if not rep:
        return None
    issue_type = gi.get("issue_type") or "grammar"
    for line_no, start, line in _line_ranges(text):
        line_end = start + len(line)
        if not (start <= offset < line_end):
            continue
        rel = offset - start
        if rel + elen > len(line):
            continue
        new_line = line[:rel] + rep + line[rel + elen :]
        stripped = line.strip()
        if _looks_like_name_only_line(stripped) or _looks_like_contact_header_line(stripped):
            return None
        if _looks_like_section_heading(stripped) and _normalize_ws(line).lower() == _normalize_ws(new_line).lower():
            return None
        if _normalize_ws(line).lower() == _normalize_ws(new_line).lower():
            return None
        if _line_has_protected_term_changed(line, new_line):
            return None
        explanation = gi.get("message", "Grammar or style suggestion.")
        if "possible spelling" in explanation.lower():
            explanation = "High-confidence spelling correction."
        return {
            "line_number": line_no,
            "issue_type": issue_type,
            "severity": "minor",
            "original_line": line,
            "improved_line": _sanitize_placeholder_text(new_line),
            "explanation": explanation,
        }
    return None


def _dedupe_line_issues(items):
    seen = set()
    out = []
    for li in items:
        if not isinstance(li, dict):
            continue
        if _is_noise_line_issue(li):
            continue
        original = li.get("original_line", "")
        improved = li.get("improved_line", "") or _first_replacement_str(li.get("replacements") or [])
        if _is_noop_line_suggestion(original, improved, li.get("rule_id", "")):
            continue
        key = (
            int(li.get("line_number", 0)),
            _line_issue_compare_text(original)[:120],
            str(li.get("issue_type", "")).lower(),
            _line_issue_compare_text(improved)[:120],
        )
        if key in seen:
            continue
        seen.add(key)
        out.append(li)
    return out


def _merge_grammar_line_issues(text: str, grammar_report: dict, llm_issues: list) -> list:
    """Combine LLM line issues with LanguageTool-derived rows (full issue list)."""
    merged = [x for x in llm_issues if isinstance(x, dict)]
    for gi in grammar_report.get("issues", []):
        if gi.get("source") == "heuristic":
            row = {
                "line_number": gi.get("line_number", 0),
                "issue_type": gi.get("issue_type", "grammar"),
                "severity": gi.get("severity", "minor"),
                "original_line": gi.get("original_line", ""),
                "improved_line": _first_replacement_str(gi.get("replacements") or []),
                "explanation": gi.get("message", "Grammar or style suggestion."),
                "rule_id": gi.get("rule_id", ""),
            }
        else:
            row = _build_grammar_line_issue(text, gi)
        if row and not _is_noise_line_issue(row):
            merged.append(row)
    return _dedupe_line_issues(merged)


def _finalize_line_issues(items: list) -> list:
    severity_rank = {"critical": 0, "important": 1, "minor": 2}
    issue_rank = {"grammar": 0, "clarity": 1, "formatting": 2, "tone": 3, "ats": 4}
    cleaned = []
    seen_suggestions = set()
    seen_issue_reason = set()
    candidates = sorted(
        [x for x in items if isinstance(x, dict)],
        key=lambda x: (
            int(x.get("line_number", 0) or 0),
            severity_rank.get(str(x.get("severity", "minor")).lower(), 3),
            issue_rank.get(str(x.get("issue_type", "clarity")).lower(), 9),
        )
    )

    for li in _dedupe_line_issues(candidates):
        if not isinstance(li, dict):
            continue
        original = li.get("original_line", "")
        improved = li.get("improved_line", "")
        rule_id = li.get("rule_id", "")
        if _is_noop_line_suggestion(original, improved, rule_id):
            continue

        original_key = _line_issue_compare_text(original)[:140]
        improved_key = _line_issue_compare_text(improved)[:140]
        issue_type = str(li.get("issue_type", "clarity")).lower()
        reason_key = _line_issue_compare_text(li.get("explanation", ""))[:80]

        suggestion_key = (int(li.get("line_number", 0) or 0), original_key, improved_key)
        if suggestion_key in seen_suggestions:
            continue
        seen_suggestions.add(suggestion_key)

        issue_reason_key = (int(li.get("line_number", 0) or 0), original_key, issue_type, reason_key)
        if issue_reason_key in seen_issue_reason:
            continue
        seen_issue_reason.add(issue_reason_key)

        cleaned.append(li)

    return sorted(
        cleaned,
        key=lambda x: (
            int(x.get("line_number", 0) or 0),
            severity_rank.get(str(x.get("severity", "minor")).lower(), 3),
            issue_rank.get(str(x.get("issue_type", "clarity")).lower(), 9),
        )
    )


def _filter_program_specificity(prog_spec: dict, prog_name: str, uni_name: str, program_fit: dict) -> dict:
    """Keep only keywords that match program-fit gaps; drop generic filler."""
    generic_kw = {
        "philosophy", "doctoral", "required", "please", "visit", "information",
        "students", "academics", "apply", "online",
    }
    kws_lower = {k.lower() for k in (program_fit.get("missing_keywords") or [])[:35]}
    placements = prog_spec.get("keyword_placement_suggestions") or []
    if not isinstance(placements, list):
        placements = []
    filtered = []
    for p in placements:
        if not isinstance(p, dict):
            continue
        kw = (p.get("keyword") or "").strip()
        if not kw or kw.lower() in generic_kw:
            continue
        if kw.lower() not in kws_lower:
            continue
        frag = _sanitize_placeholder_text(p.get("suggested_sentence_or_fragment") or "")
        if not frag.strip():
            continue
        filtered.append({
            "keyword": kw,
            "section": p.get("section", ""),
            "suggested_sentence_or_fragment": frag,
        })
    prog_spec["keyword_placement_suggestions"] = filtered[:12]
    notes = prog_spec.get("mismatch_notes") or []
    if isinstance(notes, list):
        prog_spec["mismatch_notes"] = [
            _sanitize_placeholder_text(n)
            for n in notes
            if isinstance(n, str) and len(n) < 500
        ][:10]
    return prog_spec


def grammar_check(text, max_issues=None):
    lt = _state["lt_tool"]
    matches = lt.check(text)
    protected_spans = _find_protected_spans(text)
    filtered_matches = []
    for m in matches:
        if _should_skip_lt_match(m, text):
            continue
        m_start = getattr(m, "offset", 0)
        m_end = m_start + _error_len(m)
        snip = text[m_start:m_end] if 0 <= m_start < len(text) else ""
        if _is_spelling_match(m) and not _should_keep_spelling_match(m, snip):
            continue
        if _overlaps_protected(m_start, m_end, protected_spans):
            continue
        if _is_protected_term_text(snip):
            # Likely product / project acronyms (KAIRO, etc.)
            continue
        filtered_matches.append(m)

    corrected = language_tool_python.utils.correct(text, filtered_matches)
    issues    = []
    lt_issue_rows = filtered_matches if max_issues is None else filtered_matches[:max_issues]
    for m in lt_issue_rows:
        offset = getattr(m, "offset", 0)
        elen = _error_len(m)
        line_no = _line_no_from_offset(text, offset)
        start = max(0, offset - 80)
        end = min(len(text), offset + elen + 80)
        original_line = ""
        if text:
            line_idx = min(max(0, line_no - 1), len(text.splitlines()) - 1) if text.splitlines() else 0
            if text.splitlines():
                original_line = text.splitlines()[line_idx]
        issues.append({
            "message":      m.message,
            "context":      m.context,
            "replacements": m.replacements[:5],
            "offset":       offset,
            "error_length": elen,
            "line_number":  line_no,
            "original_line": original_line,
            "rule_id":      _lt_rule_id(m),
            "category":     _lt_category(m),
            "issue_type":   _classify_lt_issue(m),
            "snippet": text[start:end]
        })
    heuristic_issues      = _heuristic_grammar_issues(text)
    issues.extend(heuristic_issues)
    issues                = _dedupe_line_issues(issues)
    grammar_issue_count   = _grammar_error_count(issues)
    grammar_quality_score = _grammar_score_from_count(grammar_issue_count)
    return {
        "count":                 grammar_issue_count,
        "grammar_error_count":   grammar_issue_count,
        "total_issue_count":     len(issues),
        "grammar_quality_score": grammar_quality_score,
        "issues":                issues,
        "corrected":             corrected
    }

def classify(text):
    prompt = (
        'Return ONLY JSON.\n'
        'Schema: {"doc_type":"SOP|RESUME","confidence":0.0,"reason_short":""}\n'
        f'TEXT:\n{text[:6000]}'
    )
    out = ask_llm_json(prompt, max_tokens=450)
    if "doc_type" not in out:
        out = {
            "doc_type":     "RESUME" if "experience" in text.lower() else "SOP",
            "confidence":   0.5,
            "reason_short": "fallback"
        }
    return out

def format_check(text, doc_type):
    lower      = text.lower()
    lines      = [l.strip() for l in text.splitlines() if l.strip()]
    first_para = lines[0] if lines else ""
    checks     = {}

    if doc_type == "RESUME":
        checks["resume_header"]     = bool(re.search(r"\bemail\b|@\w+|\bphone\b|\bmobile\b|\blinkedin\b|\bgithub\b", text, re.I))
        checks["resume_education"]  = bool(re.search(r"\beducation\b|\bdegree\b|\buniversity\b|\bcgpa\b|\bgpa\b", text, re.I))
        checks["resume_experience"] = bool(re.search(r"\bexperience\b|\bemployment\b|\bwork history\b", text, re.I))
        checks["resume_projects"]   = bool(re.search(r"\bprojects?\b|\bportfolio\b", text, re.I))
        checks["resume_skills"]     = bool(re.search(r"\bskills?\b|\btechnologies\b|\btools\b", text, re.I))
        checks["resume_extras"]     = bool(re.search(r"\bcertifications?\b|\bawards?\b|\bvolunteer\b|\bleadership\b", text, re.I))
        required = ["resume_header", "resume_education", "resume_experience", "resume_skills"]
    else:
        checks["sop_intro"]          = (
            "statement of purpose" in lower
            or "i am applying" in lower
            or "i wish to pursue" in lower
            or bool(re.search(r"\b(passionate|motivated|aspiring)\b", lower))
            or len(first_para) > 50
        )
        checks["sop_academics"]      = bool(re.search(r"\b(bachelor|undergraduate|degree|cgpa|gpa|university|graduation)\b", lower))
        checks["sop_experience"]     = bool(re.search(r"\b(intern|internship|project|research|work experience|worked)\b", lower))
        checks["sop_why_program"]    = bool(re.search(r"\b(program|curriculum|courses|specialization|research interests|master)\b", lower))
        checks["sop_why_university"] = bool(re.search(r"\b(university|faculty|lab|research group|department)\b", lower))
        checks["sop_goals"]          = bool(re.search(r"\b(career goals?|short[- ]term|long[- ]term|i aim to|i plan to|aspire|i hope to|my goal|future)\b", lower))
        checks["sop_conclusion"]     = bool(re.search(r"\b(in conclusion|thank you|i look forward|i would be honored|i am confident|i believe|i hope to)\b", lower))
        required = ["sop_intro", "sop_academics", "sop_experience", "sop_goals"]

    missing = [r for r in required if not checks.get(r)]
    if doc_type == "RESUME" and not checks.get("resume_projects"):
        missing.append("resume_projects")
    if doc_type == "SOP" and not checks.get("sop_conclusion"):
        missing.append("sop_conclusion")

    return {
        "is_correct_format": (len(missing) == 0),
        "missing_sections":  missing,
        "section_checks":    checks
    }

def extract_program_keywords(context, top_k=30):
    context_clean = re.sub(r"https?://\S+", " ", context)
    context_clean = re.sub(r"www\.\S+", " ", context_clean)
    words = re.findall(r"[a-zA-Z]{4,}", context_clean.lower())
    words = [w for w in words if w not in KEYWORD_STOPLIST]
    freq  = {}
    for w in words:
        freq[w] = freq.get(w, 0) + 1
    freq = {w: c for w, c in freq.items() if c >= 2}
    return [w for w, _ in sorted(freq.items(), key=lambda x: x[1], reverse=True)[:top_k]]


def _extract_document_keywords(text: str, top_k: int = 40) -> list:
    clean = re.sub(r"https?://\S+|www\.\S+|[\w.%+-]+@[\w.-]+", " ", text.lower())
    words = re.findall(r"[a-zA-Z][a-zA-Z+#.]{2,}", clean)
    words = [w.strip(".") for w in words if w not in KEYWORD_STOPLIST and len(w.strip(".")) >= 3]
    freq = {}
    for w in words:
        freq[w] = freq.get(w, 0) + 1
    return [w for w, _ in sorted(freq.items(), key=lambda x: x[1], reverse=True)[:top_k]]


def _normalize_domain_text(text: str) -> str:
    return re.sub(r"\s+", " ", (text or "").lower()).strip()


def _domain_text_terms(text: str) -> set:
    clean = _normalize_domain_text(text)
    raw_terms = re.findall(r"[a-z][a-z0-9+#.]*", clean)
    terms = {t.strip(".") for t in raw_terms if t.strip(".")}
    for m in re.finditer(r"[a-z][a-z0-9+#.]+(?:\s+[a-z][a-z0-9+#.]+){1,3}", clean):
        terms.add(m.group(0).strip())
    return terms


def _singularize_domain_term(term: str) -> str:
    term = term.lower().strip()
    if len(term) > 4 and term.endswith("ies"):
        return term[:-3] + "y"
    if len(term) > 3 and term.endswith("s") and not term.endswith("ss"):
        return term[:-1]
    return term


def _related_terms_for_keyword(keyword: str) -> set:
    key = _normalize_domain_text(keyword)
    related = {key}
    for concept, terms in DOMAIN_KEYWORD_MAPPING.items():
        normalized_terms = {_normalize_domain_text(t) for t in terms}
        if key == concept or key in normalized_terms:
            related.add(concept)
            related.update(normalized_terms)
    return {t for t in related if t}


def _partial_keyword_match(candidate: str, text_terms: set, normalized_text: str) -> bool:
    candidate = _normalize_domain_text(candidate)
    if not candidate:
        return False
    if " " in candidate:
        if re.search(rf"\b{re.escape(candidate)}\b", normalized_text):
            return True
        return all(_partial_keyword_match(part, text_terms, normalized_text) for part in candidate.split())
    singular_candidate = _singularize_domain_term(candidate)
    for term in text_terms:
        singular_term = _singularize_domain_term(term)
        if candidate == term or singular_candidate == singular_term:
            return True
        if len(singular_candidate) >= 5 and len(singular_term) >= 5:
            if singular_candidate.startswith(singular_term) or singular_term.startswith(singular_candidate):
                return True
    return bool(re.search(rf"\b{re.escape(candidate)}s?\b", normalized_text))


def _keyword_matches_domain_text(keyword: str, text_terms: set, normalized_text: str) -> bool:
    return any(
        _partial_keyword_match(candidate, text_terms, normalized_text)
        for candidate in _related_terms_for_keyword(keyword)
    )


def _matched_domain_keywords(keywords, text_terms: set, normalized_text: str) -> list:
    return sorted({
        _normalize_domain_text(keyword)
        for keyword in keywords
        if _keyword_matches_domain_text(keyword, text_terms, normalized_text)
    })


def _domain_relevance_boost(text: str, target_domain: str, matched_target: list) -> int:
    normalized_text = _normalize_domain_text(text)
    text_terms = _domain_text_terms(text)
    matched_count = len(matched_target)
    if matched_count == 0:
        return 0

    boost = 0
    project_signal = bool(re.search(r"\b(project|projects|built|developed|implemented|designed|created|deployed)\b", normalized_text))
    tool_signal = bool(re.search(r"\b(tool|tools|technology|technologies|skills|stack|framework|library|platform)\b", normalized_text))
    internship_signal = bool(re.search(r"\b(intern|internship|worked|experience|assistant|trainee)\b", normalized_text))

    domain_tool_hits = _matched_domain_keywords(DOMAIN_KEYWORDS.get(target_domain, set()), text_terms, normalized_text)
    if project_signal and domain_tool_hits:
        boost += 12
    if tool_signal and domain_tool_hits:
        boost += 10
    if internship_signal and domain_tool_hits:
        boost += 8
    if matched_count >= 6:
        boost += 8
    elif matched_count >= 3:
        boost += 5
    return min(boost, 25)


def _target_domain_from_text(prog_name: str, context: str = "") -> str:
    program_source = _normalize_domain_text(prog_name or "")
    context_source = _normalize_domain_text(context or "")
    aliases = {
        "computer_science": (
            "computer science", "cs", "software", "computing", "data science",
            "artificial intelligence", "machine learning", "cybersecurity",
            "information technology", "informatics",
        ),
        "law": (
            "law", "legal", "llb", "jd", "juris", "jurisprudence", "criminology",
            "public policy", "human rights", "constitutional",
        ),
        "business": (
            "business", "mba", "management", "finance", "marketing", "commerce",
            "accounting", "entrepreneurship", "economics", "analytics",
        ),
        "arts": (
            "arts", "art", "fine arts", "visual arts", "design", "graphic design",
            "media arts", "film", "animation", "photography", "music", "theatre",
            "theater", "humanities", "literature", "history",
        ),
    }
    program_terms = _domain_text_terms(program_source)
    for domain, terms in aliases.items():
        if any(_keyword_matches_domain_text(term, program_terms, program_source) for term in terms):
            return domain

    if len(program_source) >= 3:
        program_counts = {
            domain: len(_matched_domain_keywords(keywords, program_terms, program_source))
            for domain, keywords in DOMAIN_KEYWORDS.items()
        }
        program_best, program_best_count = max(program_counts.items(), key=lambda x: x[1])
        if program_best_count > 0:
            return program_best

    source = _normalize_domain_text(f"{program_source} {context_source}")
    source_terms = _domain_text_terms(source)
    for domain, terms in aliases.items():
        if any(_keyword_matches_domain_text(term, source_terms, source) for term in terms):
            return domain

    counts = {
        domain: len(_matched_domain_keywords(keywords, source_terms, source))
        for domain, keywords in DOMAIN_KEYWORDS.items()
    }
    best, best_count = max(counts.items(), key=lambda x: x[1])
    return best if best_count > 0 else "unknown"


def domain_alignment_check(text: str, prog_name: str, context: str = "") -> dict:
    normalized_text = _normalize_domain_text(text)
    text_terms = _domain_text_terms(normalized_text)
    doc_keywords = _extract_document_keywords(text, top_k=50)
    target_domain = _target_domain_from_text(prog_name, context)

    domain_hits = {}
    for domain, keywords in DOMAIN_KEYWORDS.items():
        hits = _matched_domain_keywords(keywords, text_terms, normalized_text)
        domain_hits[domain] = hits

    document_domain = "unknown"
    if domain_hits:
        best_domain, best_hits = max(domain_hits.items(), key=lambda x: len(x[1]))
        if len(best_hits) > 0:
            document_domain = best_domain

    if target_domain == "unknown":
        return {
            "target_domain": "Unknown",
            "document_domain": DOMAIN_LABELS.get(document_domain, "Unknown"),
            "score": 50,
            "is_mismatch": False,
            "message": "Target field was not specific enough for a domain alignment check.",
            "document_keywords": doc_keywords[:25],
            "target_keywords": [],
            "matched_target_keywords": [],
            "missing_target_keywords": [],
            "domain_hits": {DOMAIN_LABELS.get(k, k): v[:12] for k, v in domain_hits.items()},
        }

    target_keywords = DOMAIN_KEYWORDS[target_domain]
    matched_target = _matched_domain_keywords(target_keywords, text_terms, normalized_text)
    missing_target = sorted(target_keywords - set(matched_target))[:15]
    target_hit_count = len(matched_target)
    other_counts = {
        domain: len(hits)
        for domain, hits in domain_hits.items()
        if domain != target_domain
    }
    strongest_other_domain, strongest_other_count = max(other_counts.items(), key=lambda x: x[1]) if other_counts else ("unknown", 0)

    relevance_boost = _domain_relevance_boost(normalized_text, target_domain, matched_target)
    coverage_base = max(10, min(len(target_keywords), 20))
    coverage_score = min(100, int((target_hit_count / coverage_base) * 100))
    contrast_score = 85
    if strongest_other_count > 0:
        contrast_score = int((target_hit_count / max(1, target_hit_count + strongest_other_count)) * 100)
    score = int(round(coverage_score * 0.7 + contrast_score * 0.3)) + relevance_boost

    is_mismatch = (
        target_hit_count <= 2 and strongest_other_count >= 3 and relevance_boost < 12
    ) or (
        strongest_other_count >= target_hit_count + 3 and target_hit_count < 7 and relevance_boost < 18
    )
    if is_mismatch:
        score = min(max(score, 25), 40)
    elif target_hit_count > 0:
        score = min(score, 95)

    message = (
        "Document is well-written but not aligned with target field"
        if is_mismatch
        else "Document content is aligned with the target field."
    )

    return {
        "target_domain": DOMAIN_LABELS.get(target_domain, target_domain),
        "document_domain": DOMAIN_LABELS.get(document_domain, "Unknown"),
        "score": max(0, min(100, score)),
        "is_mismatch": is_mismatch,
        "message": message,
        "document_keywords": doc_keywords[:25],
        "target_keywords": sorted(target_keywords)[:25],
        "matched_target_keywords": matched_target[:15],
        "missing_target_keywords": missing_target,
        "domain_hits": {DOMAIN_LABELS.get(k, k): v[:12] for k, v in domain_hits.items()},
    }


def detect_tone(text: str) -> dict:
    normalized_text = _normalize_domain_text(text)
    informal_hits = []
    for phrase in INFORMAL_WORD_REPLACEMENTS:
        if re.search(rf"\b{re.escape(phrase.lower())}\b", normalized_text):
            informal_hits.append(phrase)

    professional_markers = (
        "developed", "implemented", "analyzed", "managed", "coordinated",
        "researched", "designed", "led", "improved", "professional",
        "experience", "academic", "research", "internship",
    )
    professional_hits = sum(
        1 for marker in professional_markers
        if re.search(rf"\b{re.escape(marker)}\b", normalized_text)
    )

    if informal_hits:
        shown = "', '".join(informal_hits[:3])
        return {
            "tone": "Informal",
            "message": f"Tone is informal. Avoid words like '{shown}', etc.",
            "informal_words": informal_hits[:8],
        }

    if professional_hits >= 3:
        return {
            "tone": "Professional",
            "message": "Tone is professional and appropriate.",
            "informal_words": [],
        }

    return {
        "tone": "Neutral",
        "message": "Tone is neutral and can be made more specific.",
        "informal_words": [],
    }

def program_fit_score(text, context):
    tl      = text.lower()
    kws     = extract_program_keywords(context, top_k=30)
    hits    = [k for k in kws if k in tl]
    missing = [k for k in kws if k not in tl]
    score   = min(100, int(len(hits) * (100 / max(1, len(kws)))))
    return {
        "program_fit_score":  score,
        "matched_keywords":   hits[:25],
        "missing_keywords":   missing[:25],
        "keywords_used":      kws
    }

def ats_score(text, context, fmt):
    tl               = text.lower()
    program_keywords = extract_program_keywords(context, top_k=30)

    section_score  = 0
    section_score += 10 if fmt["section_checks"].get("resume_header")     else 0
    section_score += 15 if fmt["section_checks"].get("resume_education")  else 0
    section_score += 15 if fmt["section_checks"].get("resume_experience") else 0
    section_score += 10 if fmt["section_checks"].get("resume_skills")     else 0
    section_score +=  5 if fmt["section_checks"].get("resume_projects")   else 0
    section_score  = min(section_score, 50)

    verb_count   = sum(1 for v in ACTION_VERBS if v in tl)
    verb_score   = min(15, verb_count * 2)

    metrics      = len(re.findall(r"\b\d+(\.\d+)?\s*(%|x|yrs|years|months|users|clients|requests|ms|s)\b", tl))
    metric_score = min(15, metrics * 5)

    bullets      = len(re.findall(r"(\n\s*[-•*]\s+)", text))
    bullet_score = min(10, bullets)

    prog_hits     = sum(1 for k in program_keywords if k in tl)
    prog_kw_score = min(20, prog_hits * 2)

    total = int(min(100, section_score + verb_score + metric_score + bullet_score + prog_kw_score))

    suggestions = []
    if not fmt["section_checks"].get("resume_projects"):
        suggestions.append("Add a 'Projects' section with 2–3 program-relevant projects.")
    if metrics == 0:
        suggestions.append("Add measurable impact (%, time saved, scale, users) in bullets.")
    if verb_count < 4:
        suggestions.append("Start bullets with strong action verbs (Developed, Implemented, Optimized, Led).")
    if bullets < 5:
        suggestions.append("Convert paragraphs into concise bullet points (ATS prefers bullets).")
    if prog_hits < 4:
        suggestions.append("Add program-specific keywords and technical skills aligned with the selected program.")

    return {
        "ats_score":   total,
        "breakdown":   {
            "section_score":    section_score,
            "action_verbs":     {"count": verb_count,  "score": verb_score},
            "metrics":          {"count": metrics,     "score": metric_score},
            "bullets":          {"count": bullets,     "score": bullet_score},
            "program_keywords": {"hits": prog_hits,    "score": prog_kw_score, "sample": program_keywords[:12]},
        },
        "suggestions": suggestions
    }


def build_analysis_scaffold(text: str, doc_type: str):
    """
    Build a lightweight scaffold the LLM can use to anchor feedback to
    concrete line numbers and coarse sections. This is heuristic on purpose
    but stable: we never mutate the text here, only describe it.
    """
    lines_raw = text.splitlines()
    lines = []
    for idx, t in enumerate(lines_raw, start=1):
        lines.append({"line_number": idx, "text": t})

    sections = []
    n = len(lines_raw)

    if doc_type == "RESUME":
        # Simple heading-based segmentation for resumes
        headings = [
            ("HEADER", ["name", "email", "phone", "linkedin", "github"]),
            ("SUMMARY", ["summary", "objective", "profile"]),
            ("EXPERIENCE", ["experience", "employment", "work history"]),
            ("PROJECTS", ["projects", "project"]),
            ("EDUCATION", ["education", "degree", "university"]),
            ("SKILLS", ["skills", "technologies", "tools"]),
            ("ADDITIONAL", ["certifications", "awards", "volunteer", "leadership", "activities"]),
        ]

        current_section = {"id": "FULL_RESUME", "title": "Full Resume", "start_line": 1, "end_line": n or 1, "type": "resume"}
        detected = []
        lower_lines = [l.lower() for l in lines_raw]

        for sec_id, keywords in headings:
            for i, l in enumerate(lower_lines):
                if any(re.search(rf"\\b{kw}\\b", l) for kw in keywords):
                    detected.append({"id": sec_id, "title": sec_id.title().replace("_", " "), "start_line": i + 1})
                    break

        detected_sorted = sorted({d["id"]: d for d in detected}.values(), key=lambda x: x["start_line"])
        if detected_sorted:
            for i, sec in enumerate(detected_sorted):
                start = sec["start_line"]
                end = (detected_sorted[i + 1]["start_line"] - 1) if i + 1 < len(detected_sorted) else (n or start)
                sections.append({"id": sec["id"], "title": sec["title"], "start_line": start, "end_line": end, "type": "resume"})
        else:
            sections.append(current_section)
    else:
        # SOP / Personal statement: approximate into ordered narrative sections
        ordered = [
            ("INTRODUCTION", "Introduction"),
            ("ACADEMIC_BACKGROUND", "Academic Background"),
            ("EXPERIENCE_AND_PROJECTS", "Experience and Projects"),
            ("WHY_THIS_PROGRAM", "Why This Program"),
            ("WHY_THIS_UNIVERSITY", "Why This University"),
            ("CAREER_GOALS", "Career Goals"),
            ("CONCLUSION", "Conclusion"),
        ]

        if n == 0:
            sections.append({"id": "FULL_SOP", "title": "Full Statement", "start_line": 1, "end_line": 1, "type": "sop"})
        else:
            # Divide lines into contiguous blocks following the above order
            block_size = max(1, n // len(ordered))
            start = 1
            for i, (sec_id, title) in enumerate(ordered):
                if i == len(ordered) - 1:
                    end = n
                else:
                    end = min(n, start + block_size - 1)
                sections.append({"id": sec_id, "title": title, "start_line": start, "end_line": end, "type": "sop"})
                start = end + 1
                if start > n:
                    break

    return {"lines": lines, "sections": sections}

def evaluate_and_rewrite(text, corrected, doc_type, format_report, grammar_report,
                          context, program_fit, uni_name, prog_name, ats_report=None,
                          domain_alignment=None):
    if doc_type == "RESUME":
        structure_rules = (
            "STRICT RESUME RULES:\n"
            "- MUST be ATS-friendly.\n"
            "- Use EXACT headings: NAME + CONTACT, PROFESSIONAL SUMMARY, SKILLS, EXPERIENCE, EDUCATION, PROJECTS, ADDITIONAL\n"
            "- Experience & Projects MUST be bullet points.\n"
            "- Do NOT invent facts. Use [ADD DETAIL] if missing.\n"
            "- Output MUST be valid JSON only."
        )
    else:
        structure_rules = (
            "STRICT SOP RULES:\n"
            "- Use EXACT headings: Introduction, Academic Background, Experience and Projects, "
            "Why This Program, Why This University, Career Goals, Conclusion\n"
            "- No bullet lists (use paragraphs).\n"
            "- Do NOT invent facts. Use [ADD DETAIL] if missing.\n"
            "- Output MUST be valid JSON only."
        )

    ats_block = (
        f"\nATS_REPORT:\n{json.dumps(ats_report, ensure_ascii=False)}\n"
        if ats_report else ""
    )
    grammar_summary = {
        "issue_count":            grammar_report.get("grammar_error_count", grammar_report["count"]),
        "quality_score_0_to_100": grammar_report["grammar_quality_score"],
        "top_issues":             [i["message"] for i in grammar_report["issues"] if _is_real_grammar_issue(i)][:10]
    }
    domain_alignment = domain_alignment or {}

    scaffold = build_analysis_scaffold(corrected, doc_type)

    prompt = (
        f"You are an admissions reviewer and writing coach.\n"
        f"Return ONLY valid JSON. No markdown. No triple backticks.\n"
        f"TARGET UNIVERSITY: {uni_name}\n"
        f"TARGET PROGRAM: {prog_name}\n"
        f"{structure_rules}\n"
        f"UNIVERSITY+PROGRAM CONTEXT:\n{context}\n"
        f"PROGRAM_FIT:\n{json.dumps(program_fit, ensure_ascii=False)}\n"
        f"DOMAIN_ALIGNMENT:\n{json.dumps(domain_alignment, ensure_ascii=False)}\n"
        f"FORMAT_REPORT:\n{json.dumps(format_report, ensure_ascii=False)}\n"
        f"GRAMMAR_REPORT:\n{json.dumps(grammar_summary, ensure_ascii=False)}\n"
        f"ANALYSIS_SCAFFOLD:\n{json.dumps(scaffold, ensure_ascii=False)}\n"
        f"{ats_block}"
        f"DOCUMENT:\n{corrected[:12000]}\n\n"
        "Return ONLY valid JSON with the following structure (no markdown, no comments):\n"
        "{\n"
        '  \"doc_type\": \"SOP\" | \"RESUME\",\n'
        "  \"overall_score\": number (0-100),\n"
        "  \"overall_quality_label\": string,\n"
        "  \"program_specificity\": {\n"
        "    \"target_university\": string,\n"
        "    \"target_program\": string,\n"
        "    \"fit_score\": number,\n"
        "    \"domain_alignment_score\": number,\n"
        "    \"domain_alignment_message\": string,\n"
        "    \"missing_keywords_to_add\": [string],\n"
        "    \"where_to_add_keywords\": [string],\n"
        "    \"keyword_placement_suggestions\": [\n"
        "      {\"keyword\": string, \"section\": string, \"suggested_sentence_or_fragment\": string}\n"
        "    ],\n"
        "    \"mismatch_notes\": [string]\n"
        "  },\n"
        "  \"format_and_sections\": {\n"
        "    \"is_correct_format\": boolean,\n"
        "    \"missing_sections\": [string],\n"
        "    \"section_checks\": object\n"
        "  },\n"
        "  \"grammar\": {\"score\": number, \"issue_count\": number, \"top_issues\": [string]},\n"
        "  \"clarity\": {\"score\": number, \"issues\": [string]},\n"
        "  \"tone\": {\"summary\": string, \"is_appropriate\": boolean, \"issues\": [string]},\n"
        "  \"relevance_to_selected_program\": {\n"
        "    \"score\": number,\n"
        "    \"issues\": [string],\n"
        "    \"what_to_add_based_on_context\": [string]\n"
        "  },\n"
        "  \"strengths\": [string],\n"
        "  \"weaknesses\": [string],\n"
        "  \"line_issues\": [\n"
        "    {\n"
        "      \"line_number\": number,\n"
        "      \"issue_type\": \"grammar\" | \"clarity\" | \"formatting\" | \"ats\" | \"tone\" | \"weak_content\" | \"repetition\" | \"structure\" | \"program_mismatch\",\n"
        "      \"severity\": \"critical\" | \"important\" | \"minor\",\n"
        "      \"original_line\": string,\n"
        "      \"improved_line\": string,\n"
        "      \"explanation\": string\n"
        "    }\n"
        "  ],\n"
        "  \"sentence_level_improvements\": [\n"
        "    {\n"
        "      \"section\": string,\n"
        "      \"original_sentence\": string,\n"
        "      \"improved_sentence\": string,\n"
        "      \"issue_type\": string,\n"
        "      \"severity\": string,\n"
        "      \"explanation\": string\n"
        "    }\n"
        "  ],\n"
        "  \"section_analysis\": {\n"
        "    \"SECTION_ID\": {\n"
        "      \"title\": string,\n"
        "      \"what_is_good\": [string],\n"
        "      \"what_is_missing\": [string],\n"
        "      \"what_to_improve\": [string]\n"
        "    }\n"
        "  },\n"
        "  \"resume_bullet_analysis\": [\n"
        "    {\n"
        "      \"section\": string,\n"
        "      \"bullet_text\": string,\n"
        "      \"has_action_verb\": boolean,\n"
        "      \"action_verb\": string | null,\n"
        "      \"has_metric\": boolean,\n"
        "      \"metric_example\": string | null,\n"
        "      \"program_relevance_score\": number,\n"
        "      \"issues\": [string],\n"
        "      \"improved_bullet\": string\n"
        "    }\n"
        "  ],\n"
        "  \"rewrite_output\": {\"improved_document\": string},\n"
        "  \"action_plan_next_revision\": [\n"
        "    {\"item\": string, \"priority\": \"high\" | \"medium\" | \"low\"}\n"
        "  ]\n"
        "}\n"
        "Rules:\n"
        "- Do NOT give generic advice; every issue must refer to the actual text and, when possible, a specific line or section.\n"
        "- NEVER change or \"correct\" person names, emails, phone numbers, or LinkedIn URLs. Never output placeholder headers like NAME + CONTACT.\n"
        "- NEVER change or \"correct\" technical/product/proper nouns such as FastAPI, WhisperX, LangChain, Prisma, FAISS, NaSCon, ICAP, IBA, WordPress, KAIRO, or any mixed-case/all-caps/capitalized non-standard term.\n"
        "- Do NOT suggest spelling replacements unless the correction is obvious and high confidence. If unsure, omit the issue.\n"
        "- Label only true sentence mechanics as grammar. Use clarity for sentence improvement, formatting for spacing/dashes/casing, and ats for missing keywords/metrics/resume targeting.\n"
        "- Do not mix domain mismatch, clarity, style, ATS, missing metrics, or missing details into grammar.\n"
        "- If DOMAIN_ALIGNMENT says is_mismatch is true, state that clearly in program_specificity.mismatch_notes using the exact message: Document is well-written but not aligned with target field.\n"
        "- Section headings (Education, Projects, EXPERIENCE, etc.) are formatting, not grammar; do not flag capitalization-only differences as errors.\n"
        "- Do NOT use bracket placeholders like [ADD DETAIL]; give concrete, actionable wording or say what kind of detail to add (e.g., a measurable metric).\n"
        "- Use the ANALYSIS_SCAFFOLD line numbers and sections when populating line_issues and section_analysis.\n"
        "- In line_issues, never repeat the same original line/suggestion pair. Do not include a line_issue if improved_line is identical or nearly identical to original_line.\n"
        "- Tailor feedback to the TARGET UNIVERSITY and TARGET PROGRAM using PROGRAM_FIT and the provided context; do not invent facts about the university.\n"
        "- For resumes, analyze each bullet individually for action verbs, metrics, and program relevance.\n"
        "- If information is missing, suggest what to add without fabricating numbers.\n"
    )

    out = ask_llm_json(prompt, max_tokens=3200, temperature=0.0)

    if isinstance(out, dict) and "error" not in out:
        # Ensure section/format info is always present and aligned with our own checks
        fmt_section = out.setdefault("format_and_sections", {})
        fmt_section["is_correct_format"] = format_report["is_correct_format"]
        fmt_section["missing_sections"]  = format_report["missing_sections"]
        fmt_section["section_checks"]    = format_report["section_checks"]

        # Program specificity defaults and alignment
        prog_spec = out.setdefault("program_specificity", {})
        prog_spec["target_university"]       = uni_name
        prog_spec["target_program"]          = prog_name
        prog_spec["fit_score"]               = program_fit.get("program_fit_score", 0)
        prog_spec["domain_alignment_score"]  = domain_alignment.get("score", 0)
        prog_spec["domain_alignment_message"] = domain_alignment.get("message", "")
        prog_spec["domain_alignment"]        = domain_alignment
        prog_spec.setdefault("missing_keywords_to_add", program_fit.get("missing_keywords", [])[:12])
        prog_spec.setdefault("where_to_add_keywords", [])
        prog_spec.setdefault("keyword_placement_suggestions", [])
        prog_spec.setdefault("mismatch_notes", [])
        if domain_alignment.get("is_mismatch"):
            msg = "Document is well-written but not aligned with target field"
            notes = prog_spec.get("mismatch_notes")
            if not isinstance(notes, list):
                notes = []
            if msg not in notes:
                notes.insert(0, msg)
            prog_spec["mismatch_notes"] = notes

        # Grammar defaults (scores updated after line_issues merge)
        gram = out.setdefault("grammar", {})
        gram.setdefault("top_issues", grammar_summary["top_issues"])

        # Line issues normalization + merge with LanguageTool for full coverage
        allowed_issue_types = {
            "grammar", "clarity", "formatting", "ats", "tone", "weak_content",
            "repetition", "structure", "program_mismatch",
        }
        allowed_severity    = {"critical", "important", "minor"}
        protected_values = _extract_protected_values(text)
        line_issues = out.get("line_issues") or []
        normalized_line_issues = []
        if isinstance(line_issues, list):
            for li in line_issues:
                if not isinstance(li, dict):
                    continue
                if _is_llm_line_issue_junk(li, text):
                    continue
                ln = int(li.get("line_number", 0)) or 0
                if ln <= 0:
                    continue
                itype = str(li.get("issue_type", "clarity")).lower()
                if itype not in allowed_issue_types:
                    itype = "clarity"
                expl_lower = str(li.get("explanation", "")).lower()
                if itype in {"weak_content", "program_mismatch", "structure"} and re.search(r"\b(ats|keyword|metric|resume|program|relevance)\b", expl_lower):
                    itype = "ats"
                if itype == "grammar" and re.search(r"\b(clear|clarity|simplif|wordy|readability|impact)\b", expl_lower):
                    itype = "clarity"
                if itype == "grammar" and re.search(r"\b(spacing|dash|hyphen|case|capitali[sz]ation|format)\b", expl_lower):
                    itype = "formatting"
                sev = str(li.get("severity", "minor")).lower()
                if sev not in allowed_severity:
                    sev = "minor"
                original_line = li.get("original_line", "")
                improved_line = _sanitize_placeholder_text(li.get("improved_line", ""))
                expl = _sanitize_placeholder_text(li.get("explanation", ""))
                # Preserve user identifiers (names/emails/phones/addresses) if model mutates them
                for token in protected_values:
                    if token and token in original_line and token not in improved_line:
                        improved_line = original_line
                        break
                ost = original_line.strip()
                if _looks_like_section_heading(ost) and itype == "grammar":
                    itype = "formatting"
                normalized_line_issues.append({
                    "line_number": ln,
                    "issue_type": itype,
                    "severity": sev,
                    "original_line": original_line,
                    "improved_line": improved_line,
                    "explanation": expl,
                })

        llm_filtered = [x for x in normalized_line_issues if not _is_noise_line_issue(x)]
        merged_issues = _merge_grammar_line_issues(text, grammar_report, llm_filtered)
        for x in merged_issues:
            ost = (x.get("original_line") or "").strip()
            if _looks_like_section_heading(ost) and x.get("issue_type") == "grammar":
                x["issue_type"] = "formatting"
            x["improved_line"] = _sanitize_placeholder_text(x.get("improved_line", ""))
            x["explanation"] = _sanitize_placeholder_text(x.get("explanation", ""))
        merged_issues = _finalize_line_issues(merged_issues)
        out["line_issues"] = merged_issues
        grammar_error_count = _grammar_error_count(merged_issues)
        gram["issue_count"] = grammar_error_count
        gram["total_issue_count"] = len(merged_issues)
        gram["score"] = _grammar_score_from_count(grammar_error_count)

        prog_spec = _filter_program_specificity(prog_spec, prog_name, uni_name, program_fit)

        # Section analysis normalization
        section_analysis = out.get("section_analysis") or {}
        if isinstance(section_analysis, list):
            converted = {}
            for entry in section_analysis:
                if isinstance(entry, dict):
                    key = entry.get("id") or entry.get("title") or f"section_{len(converted)+1}"
                    converted[str(key)] = {
                        "title": entry.get("title", str(key)),
                        "what_is_good": entry.get("what_is_good", []),
                        "what_is_missing": entry.get("what_is_missing", []),
                        "what_to_improve": entry.get("what_to_improve", []),
                    }
            section_analysis = converted
        elif isinstance(section_analysis, dict):
            # ensure each value has the expected keys
            for k, v in list(section_analysis.items()):
                if not isinstance(v, dict):
                    section_analysis[k] = {
                        "title": str(k),
                        "what_is_good": [],
                        "what_is_missing": [],
                        "what_to_improve": [],
                    }
                else:
                    v.setdefault("title", str(k))
                    v.setdefault("what_is_good", [])
                    v.setdefault("what_is_missing", [])
                    v.setdefault("what_to_improve", [])
        else:
            section_analysis = {}
        out["section_analysis"] = section_analysis

        # Sentence-level improvements default
        sli = out.get("sentence_level_improvements")
        if not isinstance(sli, list):
            out["sentence_level_improvements"] = []
        else:
            for s in sli:
                if isinstance(s, dict):
                    s["improved_sentence"] = _sanitize_placeholder_text(s.get("improved_sentence", ""))
                    s["explanation"] = _sanitize_placeholder_text(s.get("explanation", ""))

        # Resume bullet analysis default
        if doc_type == "RESUME":
            rba = out.get("resume_bullet_analysis")
            if not isinstance(rba, list):
                out["resume_bullet_analysis"] = []
            else:
                for b in rba:
                    if isinstance(b, dict):
                        b["improved_bullet"] = _sanitize_placeholder_text(b.get("improved_bullet", ""))
                        iss = b.get("issues")
                        if isinstance(iss, list):
                            b["issues"] = [_sanitize_placeholder_text(x) if isinstance(x, str) else x for x in iss]
        else:
            out["resume_bullet_analysis"] = []

        # Action plan normalization: allow both plain strings and objects
        ap = out.get("action_plan_next_revision") or []
        normalized_ap = []
        if isinstance(ap, list):
            for item in ap:
                if isinstance(item, str):
                    normalized_ap.append({"item": _sanitize_placeholder_text(item), "priority": "high"})
                elif isinstance(item, dict):
                    txt = _sanitize_placeholder_text(item.get("item") or item.get("text") or "")
                    pr  = str(item.get("priority", "high")).lower()
                    if pr not in {"high", "medium", "low"}:
                        pr = "high"
                    if txt:
                        normalized_ap.append({"item": txt, "priority": pr})
        out["action_plan_next_revision"] = normalized_ap

        # Derive quality label if missing
        if domain_alignment.get("is_mismatch"):
            current_score = int(out.get("overall_score", 0) or 0)
            out["overall_score"] = min(current_score if current_score else 72, 72)
        score = out.get("overall_score", 0)
        out["overall_quality_label"] = out.get("overall_quality_label") or (
            "Exceptional" if score >= 85 else
            "Strong"      if score >= 70 else
            "Average"     if score >= 50 else
            "Weak"        if score >= 30 else
            "Very Weak"
        )
    return out


# ── ROUTES ───────────────────────────────────────────────────────

@app.post("/analyze")
@app.post("/api/document-analyzer/analyze")
async def analyze_document(
    file:       UploadFile = File(...),
    university: str        = Form(...),
    program:    str        = Form(...),
):
    university = (university or "").strip()
    program = (program or "").strip()
    if not university:
        return JSONResponse({"error": "Please enter a university name."}, status_code=400)
    if not program:
        return JSONResponse({"error": "Please enter a program name."}, status_code=400)
    if not file or not file.filename:
        return JSONResponse({"error": "Please upload a document."}, status_code=400)

    suffix = os.path.splitext(file.filename)[1].lower()
    if suffix not in {".pdf", ".docx", ".txt"}:
        return JSONResponse({"error": "Unsupported file type. Please upload a PDF, DOCX, or TXT file."}, status_code=400)

    with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
        shutil.copyfileobj(file.file, tmp)
        tmp_path = tmp.name

    try:
        text = read_document(tmp_path)
        if not text.strip():
            raise ValueError("The uploaded document appears to be empty.")

        unis       = _state["unis"]
        progs      = _state["progs"]
        uni_index  = _state["uni_index"]
        prog_index = _state["prog_index"]
        embedder   = _state["embedder"]

        q = f"{university} {program}".strip()
        grammar_task = asyncio.to_thread(grammar_check, text)
        classification_task = asyncio.to_thread(classify, text)
        context_task = asyncio.to_thread(
            retrieve_context, unis, progs, uni_index, prog_index, q, embedder, 3
        )
        grammar, classification, context = await asyncio.gather(
            grammar_task, classification_task, context_task
        )
        doc_type = classification.get("doc_type", "SOP")
        fmt = format_check(text, doc_type)

        prog_fit_task = asyncio.to_thread(program_fit_score, text, context)
        domain_task = asyncio.to_thread(domain_alignment_check, text, program, context)
        tone_task = asyncio.to_thread(detect_tone, text)
        if doc_type == "RESUME":
            ats_task = asyncio.to_thread(ats_score, text, context, fmt)
            prog_fit, domain_alignment, tone_report, ats = await asyncio.gather(
                prog_fit_task, domain_task, tone_task, ats_task
            )
        else:
            prog_fit, domain_alignment, tone_report = await asyncio.gather(
                prog_fit_task, domain_task, tone_task
            )
            ats = None

        evaluation = evaluate_and_rewrite(
            text=text, corrected=grammar["corrected"], doc_type=doc_type,
            format_report=fmt, grammar_report=grammar, context=context,
            program_fit=prog_fit, uni_name=university, prog_name=program,
            ats_report=ats, domain_alignment=domain_alignment
        )

        merged_line_issues = evaluation.get("line_issues", []) if isinstance(evaluation, dict) else []
        n_grammar = _grammar_error_count(merged_line_issues)
        grammar_out = dict(grammar)
        grammar_out["count"] = n_grammar
        grammar_out["grammar_error_count"] = n_grammar
        grammar_out["total_issue_count"] = len(merged_line_issues)
        grammar_out["grammar_quality_score"] = _grammar_score_from_count(n_grammar)

        return JSONResponse({
            "classification": classification,
            "grammar":        grammar_out,
            "format":         fmt,
            "program_fit":    prog_fit,
            "domain_alignment": domain_alignment,
            "tone_detection": tone_report,
            "ats_score":      ats,
            "evaluation":     evaluation
        })

    except ValueError as ve:
        return JSONResponse({"error": str(ve)}, status_code=400)
    except Exception as e:
        import traceback
        traceback.print_exc()
        return JSONResponse({"error": f"Internal server error: {str(e)}"}, status_code=500)
    finally:
        if os.path.exists(tmp_path):
            os.unlink(tmp_path)


@app.get("/health")
@app.get("/api/document-analyzer/health")
async def health():
    return {"status": "ok"}
